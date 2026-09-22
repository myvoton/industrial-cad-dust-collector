# -*- coding: utf-8 -*-
"""交付输出: 11 张 PNG 合并 PDF + 材料明细表 xlsx"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.image as mpimg
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from dev_params import OUT, MATERIAL_ROWS

NAMES = [
    "FDC_50000_00_图纸目录及技术说明", "FDC_50000_01_总装配图",
    "FDC_50000_02_壳体部件图", "FDC_50000_03_灰斗部件图",
    "FDC_50000_04_花板零件图", "FDC_50000_05_喷吹装置部件图",
    "FDC_50000_06_顶盖及检修门部件图", "FDC_50000_07_进出风道及法兰部件图",
    "FDC_50000_08_支架平台及梯子图", "FDC_50000_09_钢板下料套裁图",
    "FDC_50000_10_材料明细表",
]


def make_pdf():
    pdf_path = os.path.join(OUT, "50000m3h布袋除尘器_全套图纸.pdf")
    with PdfPages(pdf_path) as pdf:
        for n in NAMES:
            img = mpimg.imread(os.path.join(OUT, n + ".png"))
            fig = plt.figure(figsize=(841 / 25.4, 594 / 25.4), dpi=72)
            ax = fig.add_axes([0, 0, 1, 1])
            ax.imshow(img)
            ax.axis("off")
            pdf.savefig(fig)
            plt.close(fig)
    print("PDF:", pdf_path, f"{os.path.getsize(pdf_path)/1e6:.1f}MB")


def make_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "材料明细表"
    thin = Border(*[Side(style="thin")] * 4)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    title_font = Font(name="微软雅黑", size=14, bold=True)
    head_font = Font(name="微软雅黑", size=10, bold=True)
    body_font = Font(name="微软雅黑", size=10)
    ws.merge_cells("A1:G1")
    ws["A1"] = "50000m3/h 分室脉冲喷吹袋式除尘器  材料明细表"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    header = ["序号", "名称", "规格", "数量", "材质", "单重", "备注"]
    for j, h in enumerate(header, 1):
        c = ws.cell(2, j, h)
        c.font = head_font
        c.fill = head_fill
        c.border = thin
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(MATERIAL_ROWS, 3):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = body_font
            c.border = thin
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
    for col, w in zip("ABCDEFG", (8, 16, 26, 14, 10, 10, 20)):
        ws.column_dimensions[col].width = w
    xlsx_path = os.path.join(OUT, "50000m3h布袋除尘器_材料明细表.xlsx")
    wb.save(xlsx_path)
    print("XLSX:", xlsx_path)


if __name__ == "__main__":
    make_pdf()
    make_xlsx()
